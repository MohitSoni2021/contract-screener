import io
from dataclasses import dataclass
from pathlib import Path
from typing import Union

import fitz
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

SourceType = Union[Path, str, bytes]


@dataclass(frozen=True)
class SourceBlock:
    text: str
    block_number: int
    page_number: int | None


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    blocks: list[SourceBlock]
    page_count: int | None


def _join_blocks(blocks: list[SourceBlock]) -> str:
    return "\n".join(block.text for block in blocks)


def _open_pdf(source: SourceType) -> fitz.Document:
    if isinstance(source, bytes):
        return fitz.open(stream=source, filetype="pdf")
    return fitz.open(source)


def _open_docx(source: SourceType) -> Document:
    if isinstance(source, bytes):
        return Document(io.BytesIO(source))
    return Document(source)


def pdf_page_count(source: SourceType) -> int:
    with _open_pdf(source) as document:
        return document.page_count


def _check_text_limits(
    characters: int,
    utf8_bytes: int,
    addition: str,
    max_characters: int | None,
    max_bytes: int | None,
) -> tuple[int, int]:
    characters += len(addition)
    utf8_bytes += len(addition.encode("utf-8"))
    if max_characters is not None and characters > max_characters:
        raise ValueError(
            f"This document contains more text than the configured {max_characters:,} character limit."
        )
    if max_bytes is not None and utf8_bytes > max_bytes:
        raise ValueError("The extracted text is larger than the configured text storage limit.")
    return characters, utf8_bytes


def extract_pdf(
    source: SourceType,
    max_characters: int | None = None,
    max_bytes: int | None = None,
) -> ExtractedDocument:
    source_blocks: list[SourceBlock] = []
    character_count = 0
    utf8_byte_count = 0
    with _open_pdf(source) as document:
        page_count = document.page_count
        for page_number, page in enumerate(document, start=1):
            extracted = page.get_text("blocks", sort=True)
            for item in extracted:
                if len(item) < 7 or not isinstance(item[4], str):
                    continue
                text = item[4].strip()
                if text:
                    if source_blocks:
                        character_count, utf8_byte_count = _check_text_limits(
                            character_count, utf8_byte_count, "\n", max_characters, max_bytes
                        )
                    character_count, utf8_byte_count = _check_text_limits(
                        character_count, utf8_byte_count, text, max_characters, max_bytes
                    )
                    source_blocks.append(
                        SourceBlock(
                            text=text,
                            block_number=int(item[5]),
                            page_number=page_number,
                        )
                    )
    return ExtractedDocument(_join_blocks(source_blocks), source_blocks, page_count)


def _table_text(table: Table) -> str:
    rows = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def extract_docx(
    source: SourceType,
    max_characters: int | None = None,
    max_bytes: int | None = None,
) -> ExtractedDocument:
    document = _open_docx(source)
    source_blocks: list[SourceBlock] = []
    character_count = 0
    utf8_byte_count = 0

    for block_number, element in enumerate(document.iter_inner_content()):
        if isinstance(element, Paragraph):
            text = element.text.strip()
        elif isinstance(element, Table):
            text = _table_text(element).strip()
        else:
            continue
        if text:
            if source_blocks:
                character_count, utf8_byte_count = _check_text_limits(
                    character_count, utf8_byte_count, "\n", max_characters, max_bytes
                )
            character_count, utf8_byte_count = _check_text_limits(
                character_count, utf8_byte_count, text, max_characters, max_bytes
            )
            source_blocks.append(
                SourceBlock(text=text, block_number=block_number, page_number=None)
            )

    return ExtractedDocument(_join_blocks(source_blocks), source_blocks, None)


def extract_document(
    source: SourceType,
    extension: str,
    max_characters: int | None = None,
    max_bytes: int | None = None,
) -> ExtractedDocument:
    if extension == ".pdf":
        return extract_pdf(source, max_characters=max_characters, max_bytes=max_bytes)
    if extension == ".docx":
        return extract_docx(source, max_characters=max_characters, max_bytes=max_bytes)
    raise ValueError("Only PDF and DOCX files are supported.")

